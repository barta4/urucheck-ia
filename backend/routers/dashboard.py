import logging
import io
import csv
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo
from typing import Optional, Tuple, List, Dict, Any
from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from auth import get_current_admin
from database import database
from config import settings
from rules_engine import ensure_time

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/dashboard", tags=["dashboard"])


class GeofenceCreate(BaseModel):
    """Validated input for creating a company geofence."""
    name: str = Field(..., min_length=2, max_length=100)
    latitude: float = Field(..., ge=-90.0, le=90.0)
    longitude: float = Field(..., ge=-180.0, le=180.0)
    radius_meters: int = Field(default=100, ge=10, le=10000)


def _resolve_date_range(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    default_to_last_month: bool = True
) -> Tuple[Optional[date], Optional[date]]:
    """
    Parses date strings safely. If default_to_last_month is True and dates are not specified,
    defaults to the last 30 days (date.today() - 30 days up to date.today()).
    """
    d_from = None
    if date_from and str(date_from).strip():
        try:
            d_from = date.fromisoformat(str(date_from).strip())
        except Exception:
            d_from = None

    d_to = None
    if date_to and str(date_to).strip():
        try:
            d_to = date.fromisoformat(str(date_to).strip())
        except Exception:
            d_to = None

    if default_to_last_month:
        today = date.today()
        if not d_from and not d_to:
            d_to = today
            d_from = today - timedelta(days=30)
        elif d_from and not d_to:
            d_to = today
        elif d_to and not d_from:
            d_from = d_to - timedelta(days=30)

    return d_from, d_to


def _build_log_filters(
    params: dict,
    employee_id: Optional[str],
    date_from: Optional[date],
    date_to: Optional[date],
    status: Optional[str],
) -> list[str]:
    """Build WHERE clause conditions for attendance_logs queries."""
    conditions = ["al.company_id = :cid"]
    if employee_id and str(employee_id).strip():
        conditions.append("al.employee_id = :eid")
        params["eid"] = str(employee_id).strip()
    if date_from:
        conditions.append("DATE(al.timestamp) >= :date_from")
        params["date_from"] = date_from
    if date_to:
        conditions.append("DATE(al.timestamp) <= :date_to")
        params["date_to"] = date_to
    if status and str(status).strip():
        conditions.append("al.status = :status")
        params["status"] = str(status).strip()
    return conditions

@router.get("/today")
async def today_status(admin=Depends(get_current_admin)):
    company_id = admin["company_id"]
    today = date.today()
    weekday = today.weekday() + 1

    # Single query: employees with schedule today + their check-in (if any) + streak
    # Eliminates the previous N+1 pattern (3 queries per employee)
    scheduled = await database.fetch_all(
        """
        SELECT
            e.id, e.name, e.email,
            s.start_time, s.tolerance_minutes,
            ci.timestamp  AS check_in_time,
            ci.status     AS check_in_status,
            ci.latitude   AS latitude,
            ci.longitude  AS longitude,
            st.current_streak AS current_streak,
            (
                SELECT type FROM attendance_logs al2
                WHERE al2.employee_id = e.id AND al2.company_id = :cid
                  AND DATE(al2.timestamp) = :today
                ORDER BY al2.timestamp DESC LIMIT 1
            ) AS last_mark_type
        FROM employees e
        JOIN schedules s ON e.id = s.employee_id
        LEFT JOIN LATERAL (
            SELECT timestamp, status, latitude, longitude
            FROM attendance_logs
            WHERE employee_id = e.id AND company_id = :cid
              AND type = 'check_in' AND DATE(timestamp) = :today
            LIMIT 1
        ) ci ON true
        LEFT JOIN streaks st ON st.employee_id = e.id AND st.company_id = :cid
        WHERE e.active = true AND e.company_id = :cid AND :day = ANY(s.day_of_week)
        """,
        {"cid": company_id, "day": weekday, "today": today},
    )

    now = datetime.now(ZoneInfo(settings.TIMEZONE)).time()
    result = []
    for emp in scheduled:
        if emp["check_in_time"]:
            status = emp["check_in_status"]
            current_action = emp["last_mark_type"] or "check_in"
        else:
            s_time = ensure_time(emp["start_time"])
            status = "absent" if (s_time and now > s_time) else "pending"
            current_action = None

        s_time_clean = ensure_time(emp["start_time"])
        result.append({
            "employee_id": str(emp["id"]),
            "name": emp["name"],
            "email": emp["email"],
            "scheduled_start": str(s_time_clean) if s_time_clean else str(emp["start_time"]),
            "status": status,
            "check_in_time": str(emp["check_in_time"]) if emp["check_in_time"] else None,
            "latitude": emp["latitude"],
            "longitude": emp["longitude"],
            "current_action": current_action,
            "streak": emp["current_streak"] or 0,
        })

    summary = {
        "total": len(result),
        "on_time": sum(1 for r in result if r["status"] == "on_time"),
        "late": sum(1 for r in result if r["status"] == "late"),
        "absent": sum(1 for r in result if r["status"] == "absent"),
        "pending": sum(1 for r in result if r["status"] == "pending"),
    }
    return {"summary": summary, "employees": result}

@router.get("/logs/export")
async def export_logs(
    employee_id: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    status: Optional[str] = None,
    admin=Depends(get_current_admin),
):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    d_from, d_to = _resolve_date_range(date_from, date_to, default_to_last_month=True)

    company_id = admin["company_id"]
    params = {"cid": company_id}
    conditions = _build_log_filters(params, employee_id, d_from, d_to, status)
    where = " AND ".join(conditions)
    rows = await database.fetch_all(
        f"""
        SELECT al.*, e.name as employee_name, s.slot_name, g.name as geofence_name
        FROM attendance_logs al
        JOIN employees e ON al.employee_id = e.id
        LEFT JOIN schedules s ON al.schedule_id = s.id
        LEFT JOIN geofences g ON s.geofence_id = g.id
        WHERE {where}
        ORDER BY al.timestamp DESC
        """,
        params
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registros de Asistencia"

    column_headers = ["Fecha/Hora", "Empleado", "Turno / Tramo", "Ubicación", "Tipo de Marca", "Estado", "Salida Anticipada (min)", "Latitud", "Longitud", "Racha"]
    ws.append(column_headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for row in rows:
        row_dict = dict(row)
        ts = row_dict.get("timestamp")
        ts_str = ts.strftime("%Y-%m-%d %H:%M:%S") if ts and hasattr(ts, "strftime") else str(ts or "")
        emp_name = str(row_dict.get("employee_name") or "Desconocido")
        slot_name = str(row_dict.get("slot_name") or "Turno Regular")
        geo_name = str(row_dict.get("geofence_name") or "General")
        raw_type = str(row_dict.get("type") or "")
        type_str = raw_type.replace("_", " ").title() if raw_type else "Desconocido"
        raw_status = str(row_dict.get("status") or "")
        status_str = raw_status.title() if raw_status else ""
        early_mins = row_dict.get("early_minutes") or 0
        lat = str(row_dict.get("latitude") or "")
        lng = str(row_dict.get("longitude") or "")
        streak = row_dict.get("streak_day") or 0

        ws.append([
            ts_str,
            emp_name,
            slot_name,
            geo_name,
            type_str,
            status_str,
            early_mins,
            lat,
            lng,
            streak
        ])

    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_length + 2, 12)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"Registros_Asistencia_{datetime.now(ZoneInfo(settings.TIMEZONE)).strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = {'Content-Disposition': f'attachment; filename="{filename}"'}
    return StreamingResponse(
        stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers
    )

@router.get("/logs")
async def get_logs(
    employee_id: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    status: Optional[str] = None,
    limit: int = 100,
    offset: int = 0,
    admin=Depends(get_current_admin),
):
    d_from, d_to = _resolve_date_range(date_from, date_to, default_to_last_month=True)
    company_id = admin["company_id"]
    params = {"cid": company_id, "limit": limit, "offset": offset}
    conditions = _build_log_filters(params, employee_id, d_from, d_to, status)
    where = " AND ".join(conditions)
    rows = await database.fetch_all(
        f"""
        SELECT al.*, e.name as employee_name, s.slot_name, g.name as geofence_name
        FROM attendance_logs al
        JOIN employees e ON al.employee_id = e.id
        LEFT JOIN schedules s ON al.schedule_id = s.id
        LEFT JOIN geofences g ON s.geofence_id = g.id
        WHERE {where}
        ORDER BY al.timestamp DESC
        LIMIT :limit OFFSET :offset
        """,
        params
    )
    return [dict(r) for r in rows]

@router.get("/bonus-report")
async def bonus_report(
    month: str = Query(..., description="YYYY-MM format"),
    admin=Depends(get_current_admin),
):
    company_id = admin["company_id"]
    try:
        report_month = date.fromisoformat(f"{month}-01")
    except ValueError:
        return {"error": "Formato de mes inválido. Use YYYY-MM"}

    import calendar
    _, days_in_month = calendar.monthrange(report_month.year, report_month.month)
    next_month = date(
        report_month.year,
        report_month.month + 1 if report_month.month < 12 else 1,
        1,
    )

    # Single GROUP BY query replaces previous N+1 per-employee loop
    rows = await database.fetch_all(
        """
        SELECT
            e.id, e.name, e.email,
            COUNT(DISTINCT CASE WHEN al.type = 'check_in' AND al.status = 'on_time'
                                THEN DATE(al.timestamp) END) AS on_time_days
        FROM employees e
        LEFT JOIN attendance_logs al
            ON al.employee_id = e.id AND al.company_id = :cid
            AND DATE(al.timestamp) >= :start AND DATE(al.timestamp) < :end
        WHERE e.active = true AND e.company_id = :cid
        GROUP BY e.id, e.name, e.email
        ORDER BY e.name
        """,
        {"cid": company_id, "start": report_month, "end": next_month},
    )

    report = [
        {
            "employee_id": str(r["id"]),
            "employee_name": r["name"],
            "email": r["email"],
            "month": str(report_month),
            "on_time_days": r["on_time_days"] or 0,
            "required_days": settings.BONUS_STREAK_DAYS,
            "bonus_earned": (r["on_time_days"] or 0) >= settings.BONUS_STREAK_DAYS,
        }
        for r in rows
    ]
    return {"month": month, "report": report}

@router.get("/bonus-report/export")
async def export_bonus_report(
    month: str = Query(..., description="YYYY-MM format"),
    admin=Depends(get_current_admin)
):
    data = await bonus_report(month, admin)

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=[
        "employee_name", "email", "month", "on_time_days", "required_days", "bonus_earned"
    ])
    writer.writeheader()
    for row in data["report"]:
        writer.writerow({
            "employee_name": row["employee_name"],
            "email": row["email"],
            "month": row["month"],
            "on_time_days": row["on_time_days"],
            "required_days": row["required_days"],
            "bonus_earned": "SÍ" if row["bonus_earned"] else "NO"
        })

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=bonos_{month}.csv"}
    )

@router.get("/bonus-report/pdf")
async def export_bonus_report_pdf(
    month: str = Query(..., description="YYYY-MM format"),
    admin=Depends(get_current_admin)
):
    company_id = admin["company_id"]
    company_row = await database.fetch_one(
        "SELECT name FROM companies WHERE id = :cid",
        {"cid": company_id}
    )
    company_name = company_row["name"] if company_row else "UruCheck IA"

    data = await bonus_report(month, admin)
    formatted_list = [
        {
            "employee_name": r["employee_name"],
            "email": r["email"],
            "streak_achieved": r["on_time_days"],
            "bonus_earned": r["bonus_earned"]
        }
        for r in data.get("report", [])
    ]

    from pdf_service import generate_attendance_pdf
    pdf_bytes = generate_attendance_pdf(
        company_name=company_name,
        month_str=month,
        report_data=formatted_list,
        verification_url=f"{settings.FRONTEND_URL}/verify?company_id={company_id}&month={month}"
    )

    return StreamingResponse(
        io.BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=reporte_oficial_{month}.pdf"}
    )

@router.get("/geofences")
async def list_geofences(admin=Depends(get_current_admin)):
    company_id = admin["company_id"]
    rows = await database.fetch_all(
        "SELECT * FROM geofences WHERE company_id = :cid",
        {"cid": company_id}
    )
    return [dict(r) for r in rows]

@router.post("/geofences")
async def create_geofence(data: GeofenceCreate, admin=Depends(get_current_admin)):
    company_id = admin["company_id"]

    # Check geofence limit
    from resource_limiter import check_geofence_limit
    await check_geofence_limit(company_id)

    gf_id = await database.execute(
        """
        INSERT INTO geofences (company_id, name, latitude, longitude, radius_meters)
        VALUES (:cid, :name, :lat, :lng, :radius)
        RETURNING id
        """,
        {
            "cid": company_id,
            "name": data.name,
            "lat": data.latitude,
            "lng": data.longitude,
            "radius": data.radius_meters,
        },
    )
    return {"id": str(gf_id), "message": "Geocerca creada"}

@router.post("/geofences/{geofence_id}/assign/{employee_id}")
async def assign_geofence(geofence_id: str, employee_id: str, admin=Depends(get_current_admin)):
    company_id = admin["company_id"]
    # Verify geofence belongs to company
    gf = await database.fetch_one(
        "SELECT id FROM geofences WHERE id = :gid AND company_id = :cid",
        {"gid": geofence_id, "cid": company_id}
    )
    if not gf:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="Geocerca no encontrada")

    # Verify employee belongs to company
    emp = await database.fetch_one(
        "SELECT id FROM employees WHERE id = :eid AND company_id = :cid",
        {"eid": employee_id, "cid": company_id}
    )
    if not emp:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="Empleado no encontrado")

    await database.execute(
        "INSERT INTO employee_geofences VALUES (:eid, :gid) ON CONFLICT DO NOTHING",
        {"eid": employee_id, "gid": geofence_id}
    )
    return {"message": "Geocerca asignada"}

@router.delete("/geofences/{geofence_id}")
async def delete_geofence(geofence_id: str, admin=Depends(get_current_admin)):
    company_id = admin["company_id"]
    row = await database.execute(
        "DELETE FROM geofences WHERE id = :id AND company_id = :cid RETURNING id",
        {"id": geofence_id, "cid": company_id}
    )
    if not row:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="Geocerca no encontrada")
    return {"message": "Geocerca eliminada"}
